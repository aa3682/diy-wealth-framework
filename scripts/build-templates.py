"""Build the two downloadable spreadsheet templates in public/templates/.

    pip install openpyxl
    python3 scripts/build-templates.py

The workbooks hold formulas only, no cached values, so every spreadsheet
app recalculates them on open. Formulas stick to functions that Excel,
Google Sheets, Numbers and LibreOffice all support.

The math mirrors lib/: the 50/30/20 targets round each bucket on its own
like lib/cashFlow.js, and the raise capture thresholds match
lib/raiseCapture.js. Change them together.
"""

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SITE_URL = 'https://diy-wealth-framework.vercel.app'
OUT_DIR = Path(__file__).resolve().parent.parent / 'public' / 'templates'

# Site palette (D7): slate page and cards, emerald accent.
SLATE_900 = '0F172A'
SLATE_800 = '1E293B'
SLATE_600 = '475569'
EMERALD = '10B981'
EMERALD_DARK = '047857'
INPUT_FILL = 'ECFDF5'
RED = 'B91C1C'
WHITE = 'FFFFFF'

USD = '"$"#,##0;-"$"#,##0'
PCT = '0%'
MONTH = 'mmm yyyy'

TITLE_FONT = Font(size=18, bold=True, color=SLATE_900)
SUBTITLE_FONT = Font(size=11, italic=True, color=SLATE_600)
HEADER_FONT = Font(bold=True, color=WHITE)
HEADER_FILL = PatternFill('solid', fgColor=SLATE_900)
SECTION_FONT = Font(size=13, bold=True, color=EMERALD_DARK)
LABEL_FONT = Font(bold=True, color=SLATE_900)
BODY_FONT = Font(color=SLATE_900)
RESULT_FONT = Font(bold=True, color=EMERALD_DARK)
LINK_FONT = Font(color=EMERALD_DARK, underline='single')
INPUT_STYLE = PatternFill('solid', fgColor=INPUT_FILL)
THIN = Side(style='thin', color='A7F3D0')
INPUT_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')

# Fixed timestamps keep rebuilds from changing the file metadata.
BUILD_TIME = datetime(2026, 1, 1)


def new_workbook():
    wb = Workbook()
    wb.properties.creator = 'DIY Wealth Framework'
    wb.properties.created = BUILD_TIME
    wb.properties.modified = BUILD_TIME
    return wb


def title(ws, text, subtitle=None):
    ws['A1'] = text
    ws['A1'].font = TITLE_FONT
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws['A2'] = subtitle
        ws['A2'].font = SUBTITLE_FONT


def header_row(ws, row, values, start_col=1):
    for i, value in enumerate(values):
        cell = ws.cell(row=row, column=start_col + i, value=value)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center' if i else 'left', vertical='center')


def as_input(cell, number_format=None):
    cell.fill = INPUT_STYLE
    cell.border = INPUT_BORDER
    cell.font = BODY_FONT
    if number_format:
        cell.number_format = number_format


def label(ws, ref, text, font=LABEL_FONT):
    ws[ref] = text
    ws[ref].font = font


def widths(ws, **cols):
    for col, width in cols.items():
        ws.column_dimensions[col].width = width


def start_here(ws, heading, intro, steps, link_text, link_path):
    title(ws, heading, intro)
    widths(ws, A=4, B=96)
    ws.sheet_view.showGridLines = False
    row = 4
    label(ws, f'A{row}', 'How to use it', SECTION_FONT)
    for i, step in enumerate(steps, start=1):
        row += 1
        ws[f'A{row}'] = f'{i}.'
        ws[f'A{row}'].font = LABEL_FONT
        ws[f'A{row}'].alignment = Alignment(vertical='top')
        ws[f'B{row}'] = step
        ws[f'B{row}'].font = BODY_FONT
        ws[f'B{row}'].alignment = WRAP
        ws.row_dimensions[row].height = 32 if len(step) > 95 else 16
    row += 2
    label(ws, f'A{row}', 'Color key', SECTION_FONT)
    row += 1
    ws[f'B{row}'] = 'Green cells are yours to fill in. Everything else is a formula; leave it alone.'
    as_input(ws[f'B{row}'])
    row += 2
    ws[f'B{row}'] = link_text
    ws[f'B{row}'].hyperlink = f'{SITE_URL}{link_path}'
    ws[f'B{row}'].font = LINK_FONT
    row += 1
    ws[f'B{row}'] = 'Educational template only, not financial advice. No data leaves your spreadsheet.'
    ws[f'B{row}'].font = SUBTITLE_FONT


# --------------------------------------------------------------------------
# Cash Flow Tracker
# --------------------------------------------------------------------------

MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
BUCKETS = ['Needs', 'Wants', 'Future You']

CATEGORIES = [
    ('Rent / mortgage', 'Needs'),
    ('Groceries', 'Needs'),
    ('Utilities', 'Needs'),
    ('Insurance', 'Needs'),
    ('Transportation', 'Needs'),
    ('Minimum debt payments', 'Needs'),
    ('Dining out', 'Wants'),
    ('Travel', 'Wants'),
    ('Entertainment', 'Wants'),
    ('Subscriptions', 'Wants'),
    ('Shopping', 'Wants'),
    ('Investing sweep', 'Future You'),
    ('High-yield savings', 'Future You'),
    ('Extra debt paydown', 'Future You'),
]

CAT_FIRST, CAT_LAST = 16, 55  # 40 category rows


def cash_flow_setup(ws):
    title(ws, 'Setup', 'Enter your monthly take-home pay once. Every month on the Monthly tab starts from it.')
    widths(ws, A=34, B=16, C=70)
    ws.sheet_view.showGridLines = False
    label(ws, 'A5', 'Monthly take-home pay')
    as_input(ws['B5'], USD)

    label(ws, 'A7', '50/30/20 targets', SECTION_FONT)
    targets = [
        ('Needs (50%)', 0.5, 'Housing, groceries, utilities, minimum debt payments.'),
        ('Wants (30%)', 0.3, 'Dining out, travel, entertainment.'),
        ('Future You (20%)', 0.2, 'Debt paydown above minimums, investing and saving.'),
    ]
    for i, (name, share, note) in enumerate(targets):
        row = 8 + i
        label(ws, f'A{row}', name, BODY_FONT)
        ws[f'B{row}'] = f'=ROUND($B$5*{share},0)'
        ws[f'B{row}'].number_format = USD
        ws[f'B{row}'].font = RESULT_FONT
        ws[f'C{row}'] = note
        ws[f'C{row}'].font = SUBTITLE_FONT

    ws['A12'] = ('Each target is rounded to the nearest dollar on its own, so the three can add up '
                 'to $1 more or less than your pay. The calculator on the site does the same.')
    ws['A12'].font = SUBTITLE_FONT
    ws.merge_cells('A12:C12')
    ws['A12'].alignment = WRAP
    ws.row_dimensions[12].height = 30
    ws['A14'] = ('If Needs keep landing above 50%, you cannot budget your way out of it: you have a '
                 'structural housing or transportation cost problem.')
    ws['A14'].font = SUBTITLE_FONT
    ws.merge_cells('A14:C14')
    ws['A14'].alignment = WRAP
    ws.row_dimensions[14].height = 30


def cash_flow_monthly(ws):
    title(ws, 'Monthly',
          'One row per category, not per transaction. Enter what went to each category each month.')
    widths(ws, A=28, B=14, O=14)
    for col in range(3, 15):
        ws.column_dimensions[get_column_letter(col)].width = 11
    ws.freeze_panes = 'C5'

    header_row(ws, 4, ['Summary', 'Target'] + MONTHS + ['Year to date'])
    cats = f'$B${CAT_FIRST}:$B${CAT_LAST}'

    rows = {
        5: 'Take-home pay',
        6: 'Needs',
        7: 'Wants',
        8: 'Future You',
        9: 'Unassigned',
        10: 'Needs % of pay',
        11: 'Wants % of pay',
        12: 'Future You % (savings rate)',
        13: 'Needs check',
    }
    for row, text in rows.items():
        label(ws, f'A{row}', text, LABEL_FONT if row in (5, 9, 12) else BODY_FONT)

    for row, ref in ((6, 'B8'), (7, 'B9'), (8, 'B10')):
        ws[f'B{row}'] = f'=Setup!{ref}'
        ws[f'B{row}'].number_format = USD
    for row, share in ((10, 0.5), (11, 0.3), (12, 0.2)):
        ws[f'B{row}'] = share
        ws[f'B{row}'].number_format = PCT
    ws['B13'] = 'Needs ≤ 50%'
    for row in range(6, 14):
        ws[f'B{row}'].font = Font(color=SLATE_600)

    for col in range(3, 16):  # C..N months, O year
        c = get_column_letter(col)
        is_year = col == 15
        if is_year:
            # Only months with entries count, so the year's percentages
            # are not diluted by months still to come.
            ws[f'{c}5'] = '=SUMPRODUCT((C6:N6+C7:N7+C8:N8>0)*C5:N5)'
        else:
            ws[f'{c}5'] = '=Setup!$B$5'
            as_input(ws[f'{c}5'])
        ws[f'{c}5'].number_format = USD
        for row, bucket in ((6, 'Needs'), (7, 'Wants'), (8, 'Future You')):
            ws[f'{c}{row}'] = f'=SUMIF({cats},"{bucket}",{c}${CAT_FIRST}:{c}${CAT_LAST})'
            ws[f'{c}{row}'].number_format = USD
        spent = f'SUM({c}6:{c}8)'
        ws[f'{c}9'] = f'=IF({spent}=0,"",{c}5-{spent})'
        ws[f'{c}9'].number_format = USD
        ws[f'{c}9'].font = LABEL_FONT
        for row, src in ((10, 6), (11, 7), (12, 8)):
            ws[f'{c}{row}'] = f'=IF(OR({spent}=0,{c}$5<=0),"",{c}{src}/{c}$5)'
            ws[f'{c}{row}'].number_format = PCT
        ws[f'{c}12'].font = RESULT_FONT
        ws[f'{c}13'] = f'=IF({c}10="","",IF({c}10>0.5,"Over 50%","OK"))'
        ws[f'{c}13'].alignment = Alignment(horizontal='right')

    red = Font(bold=True, color=RED)
    ws.conditional_formatting.add('C10:O10', FormulaRule(formula=['AND(ISNUMBER(C10),C10>0.5)'], font=red))
    ws.conditional_formatting.add('C13:O13', FormulaRule(formula=['C13="Over 50%"'], font=red))
    ws.conditional_formatting.add('C9:O9', FormulaRule(formula=['AND(ISNUMBER(C9),C9<0)'], font=red))

    header_row(ws, CAT_FIRST - 1, ['Category', 'Bucket'] + MONTHS + ['Year'])
    bucket_list = DataValidation(type='list', formula1='"Needs,Wants,Future You"', allow_blank=True)
    ws.add_data_validation(bucket_list)
    for i, row in enumerate(range(CAT_FIRST, CAT_LAST + 1)):
        name, bucket = CATEGORIES[i] if i < len(CATEGORIES) else (None, None)
        as_input(ws[f'A{row}'])
        as_input(ws[f'B{row}'])
        ws[f'A{row}'] = name
        ws[f'B{row}'] = bucket
        bucket_list.add(ws[f'B{row}'])
        for col in range(3, 15):
            as_input(ws.cell(row=row, column=col), USD)
        ws[f'O{row}'] = f'=IF(COUNT(C{row}:N{row})=0,"",SUM(C{row}:N{row}))'
        ws[f'O{row}'].number_format = USD


def cash_flow_raise(ws):
    title(ws, 'Raise Capture',
          'How much of your last raise reached savings? Use annual figures from two points in time.')
    widths(ws, A=30, B=16, C=4, D=70)
    ws.sheet_view.showGridLines = False
    for row, text in ((5, 'Past annual income'), (6, 'Current annual income'),
                      (7, 'Past annual savings'), (8, 'Current annual savings')):
        label(ws, f'A{row}', text, BODY_FONT)
        as_input(ws[f'B{row}'], USD)

    label(ws, 'A10', 'Income increase', BODY_FONT)
    ws['B10'] = '=IF(COUNT(B5:B8)<4,"",B6-B5)'
    label(ws, 'A11', 'Savings increase', BODY_FONT)
    ws['B11'] = '=IF(COUNT(B5:B8)<4,"",B8-B7)'
    for ref in ('B10', 'B11'):
        ws[ref].number_format = USD
    label(ws, 'A12', 'Capture rate')
    ws['B12'] = '=IF(COUNT(B5:B8)<4,"",IF(B10>0,B11/B10,0))'
    ws['B12'].number_format = PCT
    ws['B12'].font = RESULT_FONT

    label(ws, 'A14', 'Verdict')
    ws['B14'] = ('=IF(COUNT(B5:B8)<4,"Enter all four numbers above",'
                 'IF(B10<=0,"Income stagnant or decreased",'
                 'IF(B12>=0.5,"Excellent wealth capture",'
                 'IF(B12>=0.2,"Moderate creep","Severe lifestyle creep"))))')
    ws['B14'].font = RESULT_FONT
    label(ws, 'A15', 'What to do')
    ws['B15'] = (
        '=IF(COUNT(B5:B8)<4,"",'
        'IF(B10<=0,"Your income has not increased during this period. Focus on increasing your earning '
        'power or reducing baseline expenses.",'
        'IF(B12>=0.5,"Your cash flow systems are successfully resisting lifestyle inflation.",'
        'IF(B12>=0.2,"You are hitting the baseline 20% target, but your lifestyle is inflating '
        'noticeably. Consider routing your next raise entirely to investments.",'
        '"You are spending almost all of your new money. You urgently need an intermediate Holding '
        'Account to trap raises before they hit your checking account."))))')
    ws['B15'].font = BODY_FONT
    ws['B15'].alignment = WRAP
    ws.merge_cells('B15:D17')

    ws.conditional_formatting.add('B14', FormulaRule(formula=['B14="Severe lifestyle creep"'],
                                                     font=Font(bold=True, color=RED)))
    ws['A19'] = 'Thresholds: 50% or more is excellent, 20–49% is moderate, under 20% is severe.'
    ws['A19'].font = SUBTITLE_FONT


def build_cash_flow_tracker():
    wb = new_workbook()
    start = wb.active
    start.title = 'Start Here'
    start_here(
        start,
        'Cash Flow Tracker',
        'A monthly 50/30/20 check-in from the DIY Wealth Framework.',
        [
            'On the Setup tab, enter your monthly take-home pay. Your 50/30/20 targets fill in.',
            'On the Monthly tab, rename, add or delete categories to fit your life and tag each one '
            'Needs, Wants or Future You. The categories provided are only examples.',
            'Once a month, enter what went to each category. Round numbers are fine: this is a '
            'check-in, not a ledger, so there is no need to track every coffee.',
            'Read the summary at the top of Monthly: your savings rate, and a red flag when Needs '
            'take more than 50% of pay. If one month\'s pay was different, overwrite it in row 5.',
            'After a raise, use the Raise Capture tab to see how much of it reached savings.',
        ],
        'Read the guide: Cash Flow Systems →',
        '/02-cash-flow-systems',
    )
    cash_flow_setup(wb.create_sheet('Setup'))
    cash_flow_monthly(wb.create_sheet('Monthly'))
    cash_flow_raise(wb.create_sheet('Raise Capture'))
    return wb


# --------------------------------------------------------------------------
# Net Worth Dashboard
# --------------------------------------------------------------------------

# Liquid marks the assets you could spend within days, and the debts that
# count against them. The mortgage and auto loan are left out because the
# home and car securing them are left out too.
TYPES = [
    ('Cash', 'Asset', 'Yes'),
    ('Taxable investing', 'Asset', 'Yes'),
    ('Retirement', 'Asset', 'No'),
    ('HSA', 'Asset', 'No'),
    ('Real estate', 'Asset', 'No'),
    ('Vehicle', 'Asset', 'No'),
    ('Other asset', 'Asset', 'No'),
    ('Mortgage', 'Debt', ''),
    ('Student loan', 'Debt', 'Yes'),
    ('Auto loan', 'Debt', ''),
    ('Credit card', 'Debt', 'Yes'),
    ('Other debt', 'Debt', 'Yes'),
]
ASSET_TYPES = [t for t, kind, _ in TYPES if kind == 'Asset']

ACCOUNTS = [
    ('Checking', 'Cash'),
    ('High-yield savings', 'Cash'),
    ('Brokerage', 'Taxable investing'),
    ('401(k)', 'Retirement'),
    ('Roth IRA', 'Retirement'),
    ('HSA', 'HSA'),
    ('Home', 'Real estate'),
    ('Car', 'Vehicle'),
    ('Mortgage', 'Mortgage'),
    ('Auto loan', 'Auto loan'),
    ('Student loan', 'Student loan'),
    ('Credit card', 'Credit card'),
]

ACC_FIRST, ACC_LAST = 5, 44  # 40 account rows
N_MONTHS = 24
FIRST_MONTH_COL = 3  # C
LAST_MONTH = get_column_letter(FIRST_MONTH_COL + N_MONTHS - 1)  # Z
TOTAL_ROWS = {'assets': 47, 'debts': 48, 'net': 49, 'liquid': 50, 'liquid_net': 51, 'cash': 52, 'has': 53}


def net_worth_lists(ws):
    title(ws, 'Lists', 'Account types used by the Accounts tab. Edit with care.')
    widths(ws, A=20, B=10, C=10)
    header_row(ws, 4, ['Type', 'Kind', 'Liquid'])
    for i, row_values in enumerate(TYPES):
        for j, value in enumerate(row_values):
            ws.cell(row=5 + i, column=1 + j, value=value or None).font = BODY_FONT


def net_worth_accounts(ws):
    types_range = f'Lists!$A$5:$C${4 + len(TYPES)}'
    title(ws, 'Accounts', 'List every asset and debt once. Balances go on the Snapshots tab.')
    widths(ws, A=28, B=20, C=10, D=10, E=36)
    ws.freeze_panes = 'A5'
    header_row(ws, 4, ['Account', 'Type', 'Kind', 'Liquid', 'Notes'])
    type_list = DataValidation(type='list', formula1=f'Lists!$A$5:$A${4 + len(TYPES)}', allow_blank=True)
    ws.add_data_validation(type_list)
    for i, row in enumerate(range(ACC_FIRST, ACC_LAST + 1)):
        name, kind = ACCOUNTS[i] if i < len(ACCOUNTS) else (None, None)
        for col in 'ABE':
            as_input(ws[f'{col}{row}'])
        ws[f'A{row}'] = name
        ws[f'B{row}'] = kind
        type_list.add(ws[f'B{row}'])
        ws[f'C{row}'] = f'=IF(B{row}="","",VLOOKUP(B{row},{types_range},2,FALSE))'
        ws[f'D{row}'] = f'=IF(B{row}="","",IF(VLOOKUP(B{row},{types_range},3,FALSE)="Yes","Yes",""))'
        for col in 'CD':
            ws[f'{col}{row}'].font = Font(color=SLATE_600)
    ws[f'A{ACC_LAST + 2}'] = 'Enter debts as positive balances. The Kind column subtracts them for you.'
    ws[f'A{ACC_LAST + 2}'].font = SUBTITLE_FONT


def net_worth_snapshots(ws):
    title(ws, 'Snapshots',
          'Once a month, enter each balance. Set your first month in C4; the rest follow.')
    widths(ws, A=28, B=18)
    for col in range(FIRST_MONTH_COL, FIRST_MONTH_COL + N_MONTHS):
        ws.column_dimensions[get_column_letter(col)].width = 12
    ws.freeze_panes = 'C5'
    header_row(ws, 4, ['Account', 'Type'])
    for i in range(N_MONTHS):
        col = get_column_letter(FIRST_MONTH_COL + i)
        cell = ws[f'{col}4']
        if i == 0:
            cell.value = date(2026, 1, 1)
            as_input(cell)
            cell.font = Font(bold=True, color=SLATE_900)
        else:
            prev = get_column_letter(FIRST_MONTH_COL + i - 1)
            cell.value = f'=EDATE({prev}4,1)'
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
        cell.number_format = MONTH
        cell.alignment = Alignment(horizontal='center')

    for row in range(ACC_FIRST, ACC_LAST + 1):
        ws[f'A{row}'] = f'=IF(Accounts!A{row}="","",Accounts!A{row})'
        ws[f'B{row}'] = f'=IF(Accounts!B{row}="","",Accounts!B{row})'
        ws[f'B{row}'].font = Font(color=SLATE_600)
        for col in range(FIRST_MONTH_COL, FIRST_MONTH_COL + N_MONTHS):
            as_input(ws.cell(row=row, column=col), USD)

    header_row(ws, TOTAL_ROWS['assets'] - 1, ['Totals', ''] + [None] * N_MONTHS)
    labels = {
        'assets': 'Total assets',
        'debts': 'Total debts',
        'net': 'Net worth',
        'liquid': 'Liquid assets',
        'liquid_net': 'Liquid net worth',
        'cash': 'Cash',
        'has': 'Month has balances',
    }
    for key, row in TOTAL_ROWS.items():
        label(ws, f'A{row}', labels[key], LABEL_FONT if key in ('net', 'liquid_net') else BODY_FONT)
    ws[f'B{TOTAL_ROWS["has"]}'] = '(1 = yes; used by Dashboard)'
    ws[f'B{TOTAL_ROWS["has"]}'].font = SUBTITLE_FONT

    for i in range(N_MONTHS):
        c = get_column_letter(FIRST_MONTH_COL + i)
        vals = f'{c}${ACC_FIRST}:{c}${ACC_LAST}'
        has = f'{c}${TOTAL_ROWS["has"]}'
        kind = f'Accounts!$C${ACC_FIRST}:$C${ACC_LAST}'
        liquid = f'Accounts!$D${ACC_FIRST}:$D${ACC_LAST}'
        formulas = {
            'assets': f'SUMIFS({vals},Accounts!$C${ACC_FIRST}:$C${ACC_LAST},"Asset")',
            'debts': f'SUMIFS({vals},Accounts!$C${ACC_FIRST}:$C${ACC_LAST},"Debt")',
            'net': f'{c}{TOTAL_ROWS["assets"]}-{c}{TOTAL_ROWS["debts"]}',
            'liquid': f'SUMIFS({vals},{kind},"Asset",{liquid},"Yes")',
            'liquid_net': f'{c}{TOTAL_ROWS["liquid"]}-SUMIFS({vals},{kind},"Debt",{liquid},"Yes")',
            'cash': f'SUMIFS({vals},Accounts!$B${ACC_FIRST}:$B${ACC_LAST},"Cash")',
        }
        for key, formula in formulas.items():
            cell = ws[f'{c}{TOTAL_ROWS[key]}']
            cell.value = f'=IF({has}=1,{formula},"")'
            cell.number_format = USD
            cell.font = RESULT_FONT if key in ('net', 'liquid_net') else BODY_FONT
        ws[f'{c}{TOTAL_ROWS["has"]}'] = f'=IF(COUNT({vals})>0,1,"")'
        ws[f'{c}{TOTAL_ROWS["has"]}'].font = Font(color=SLATE_600)
        ws[f'{c}{TOTAL_ROWS["has"]}'].alignment = Alignment(horizontal='center')


def net_worth_dashboard(ws):
    title(ws, 'Dashboard', 'Your latest month at a glance. Nothing to enter here except monthly essentials.')
    widths(ws, A=30, B=16, C=10, D=4, E=14, F=16)
    ws.sheet_view.showGridLines = False

    months = f'Snapshots!$C$4:${LAST_MONTH}$4'
    has = f'Snapshots!$C${TOTAL_ROWS["has"]}:${LAST_MONTH}${TOTAL_ROWS["has"]}'
    latest = f'MATCH(2,{has})'  # last month with balances

    def row_range(key):
        return f'Snapshots!$C${TOTAL_ROWS[key]}:${LAST_MONTH}${TOTAL_ROWS[key]}'

    def at_latest(key):
        return f'=IFERROR(INDEX({row_range(key)},{latest}),"")'

    label(ws, 'A4', 'Latest month')
    ws['B4'] = f'=IFERROR(INDEX({months},{latest}),"No balances yet")'
    ws['B4'].number_format = MONTH
    ws['B4'].font = RESULT_FONT
    ws['B4'].alignment = Alignment(horizontal='right')

    label(ws, 'A6', 'Where you stand', SECTION_FONT)
    for row, text, key in ((7, 'Net worth', 'net'), (8, 'Total assets', 'assets'),
                           (9, 'Total debts', 'debts'), (10, 'Liquid net worth', 'liquid_net')):
        label(ws, f'A{row}', text, LABEL_FONT if key in ('net', 'liquid_net') else BODY_FONT)
        ws[f'B{row}'] = at_latest(key)
        ws[f'B{row}'].number_format = USD
        ws[f'B{row}'].font = RESULT_FONT if key in ('net', 'liquid_net') else BODY_FONT
    label(ws, 'A11', 'Change since previous month', BODY_FONT)
    net = row_range('net')
    ws['B11'] = (f'=IFERROR(IF(INDEX({has},{latest}-1)=1,'
                 f'B7-INDEX({net},{latest}-1),""),"")')
    ws['B11'].number_format = '+"$"#,##0;-"$"#,##0;"$"0'
    label(ws, 'A12', 'Debt-to-asset ratio', BODY_FONT)
    ws['B12'] = '=IF(AND(ISNUMBER(B8),B8>0),B9/B8,"")'
    ws['B12'].number_format = PCT
    ws['A13'] = ('Liquid net worth: cash and taxable investing, minus debts other than the '
                 'mortgage and auto loan.')
    ws['A13'].font = SUBTITLE_FONT
    ws['A13'].alignment = WRAP
    ws.merge_cells('A13:C14')

    label(ws, 'A15', 'Emergency reserve', SECTION_FONT)
    label(ws, 'A16', 'Monthly essential expenses', BODY_FONT)
    as_input(ws['B16'], USD)
    label(ws, 'A17', 'Cash on hand', BODY_FONT)
    ws['B17'] = at_latest('cash')
    ws['B17'].number_format = USD
    label(ws, 'A18', 'Months of expenses covered')
    ws['B18'] = '=IF(AND(ISNUMBER(B16),B16>0,ISNUMBER(B17)),B17/B16,"")'
    ws['B18'].number_format = '0.0" months"'
    ws['B18'].font = RESULT_FONT
    ws['A19'] = 'Size your target with the Emergency Reserve calculator in section 1.'
    ws['A19'].hyperlink = f'{SITE_URL}/01-financial-defense'
    ws['A19'].font = LINK_FONT

    mix_first = 23
    label(ws, f'A{mix_first - 2}', 'Asset mix (latest month)', SECTION_FONT)
    header_row(ws, mix_first - 1, ['Type', 'Balance', 'Share'])
    month_values = f'INDEX(Snapshots!$C${ACC_FIRST}:${LAST_MONTH}${ACC_LAST},0,{latest})'
    for i, asset_type in enumerate(ASSET_TYPES):
        row = mix_first + i
        ws[f'A{row}'] = asset_type
        ws[f'A{row}'].font = BODY_FONT
        ws[f'B{row}'] = (f'=IFERROR(SUMIFS({month_values},'
                         f'Accounts!$B${ACC_FIRST}:$B${ACC_LAST},A{row}),"")')
        ws[f'B{row}'].number_format = USD
        ws[f'C{row}'] = f'=IF(AND(ISNUMBER(B{row}),ISNUMBER($B$8),$B$8>0),B{row}/$B$8,"")'
        ws[f'C{row}'].number_format = PCT
    mix_last = mix_first + len(ASSET_TYPES) - 1

    # Trend table feeding the chart. Empty months are #N/A, which Excel and
    # Google Sheets leave out of a line chart (LibreOffice draws them as $0);
    # conditional formatting hides the #N/A text. A column chart avoids the
    # LibreOffice drop but exaggerates small changes, since its axis does not
    # start at $0 when values are close together.
    header_row(ws, 6, ['Month', 'Net worth'], start_col=5)
    for i in range(N_MONTHS):
        row = 7 + i
        n = i + 1
        ws[f'E{row}'] = f'=INDEX({months},{n})'
        ws[f'E{row}'].number_format = MONTH
        ws[f'E{row}'].font = BODY_FONT
        ws[f'F{row}'] = f'=IF(INDEX({has},{n})=1,INDEX({net},{n}),NA())'
        ws[f'F{row}'].number_format = USD
        ws[f'F{row}'].font = BODY_FONT
    trend_last = 6 + N_MONTHS
    ws.conditional_formatting.add(f'F7:F{trend_last}',
                                  FormulaRule(formula=['ISERROR(F7)'], font=Font(color=WHITE)))

    trend = LineChart()
    trend.title = 'Net worth by month'
    trend.style = 2
    trend.height, trend.width = 7.5, 16
    trend.legend = None
    trend.display_blanks = 'gap'
    trend.y_axis.number_format = '"$"#,##0'
    trend.x_axis.number_format = 'mmm yy'
    trend.add_data(Reference(ws, min_col=6, min_row=6, max_row=trend_last), titles_from_data=True)
    trend.set_categories(Reference(ws, min_col=5, min_row=7, max_row=trend_last))
    series = trend.series[0]
    series.graphicalProperties.line.solidFill = EMERALD
    series.graphicalProperties.line.width = 28000
    series.smooth = False
    ws.add_chart(trend, 'H4')

    mix = BarChart()
    mix.type = 'bar'
    mix.title = 'Asset mix'
    mix.style = 2
    mix.height, mix.width = 7.5, 16
    mix.legend = None
    mix.x_axis.number_format = '"$"#,##0'
    mix.add_data(Reference(ws, min_col=2, min_row=mix_first - 1, max_row=mix_last), titles_from_data=True)
    mix.set_categories(Reference(ws, min_col=1, min_row=mix_first, max_row=mix_last))
    mix.series[0].graphicalProperties.solidFill = SLATE_800
    ws.add_chart(mix, 'H21')


def build_net_worth_dashboard():
    wb = new_workbook()
    start = wb.active
    start.title = 'Start Here'
    start_here(
        start,
        'Net Worth Dashboard',
        'A monthly snapshot of everything you own and owe, from the DIY Wealth Framework.',
        [
            'On the Accounts tab, list every asset and debt once and pick its type. The accounts '
            'provided are only examples: rename, add or delete rows.',
            'Once a month, go to the Snapshots tab and enter each balance in that month\'s column. '
            'Enter debts as positive numbers.',
            'Set your first month in Snapshots cell C4. The next 23 months fill in after it.',
            'On the Dashboard tab, enter your monthly essential expenses once to see how many months '
            'your cash covers.',
            'Read the Dashboard: net worth, liquid net worth, debt-to-asset ratio, asset mix and '
            'the trend. It always shows the latest month you filled in.',
        ],
        'Read the guide: Financial Defense →',
        '/01-financial-defense',
    )
    net_worth_dashboard(wb.create_sheet('Dashboard'))
    net_worth_accounts(wb.create_sheet('Accounts'))
    net_worth_snapshots(wb.create_sheet('Snapshots'))
    net_worth_lists(wb.create_sheet('Lists'))
    return wb


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for name, build in (('cash-flow-tracker.xlsx', build_cash_flow_tracker),
                        ('net-worth-dashboard.xlsx', build_net_worth_dashboard)):
        path = OUT_DIR / name
        wb = build()
        for ws in wb:
            ws.page_setup.orientation = 'landscape'
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
        wb.save(path)
        print(f'wrote {path.relative_to(OUT_DIR.parent.parent)}')


if __name__ == '__main__':
    main()
